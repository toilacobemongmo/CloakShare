#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Kiểm tra các công thức trong DOCX đầu ra: còn sót $ hay backslash không."""
import re
import sys
import zipfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
DOCX = ROOT / "CloakShare_Bao_Cao_Chuyen_Khao_400_Trang.docx"

with zipfile.ZipFile(DOCX) as z:
    xml = z.read("word/document.xml").decode("utf-8")

texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, re.DOTALL)
texts = [re.sub(r"<[^>]+>", "", t) for t in texts]
full = "".join(texts)

print("=== KIEM TRA CONG THUC TRONG DOCX ===")
checks = {
    "Ky tu $ (dau $$ con sot)": full.count("$"),
    "Lenh \\text": full.count("\\text"),
    "Lenh \\mathcal": full.count("\\mathcal"),
    "Lenh \\oplus": full.count("\\oplus"),
    "Lenh \\begin / \\end": full.count("\\begin") + full.count("\\end"),
    "Lenh \\frac / \\sqrt": full.count("\\frac") + full.count("\\sqrt"),
    "Escape \\= / \\> / \\<": len(re.findall(r"\\[=<>]", full)),
}
for k, v in checks.items():
    flag = "OK " if v == 0 else "CON"
    print(f"  [{flag}] {k:<34} {v}")

print()
print("=== KY TU TOAN HOC DA DICH (dau thu) ===")
uni = {c: full.count(c) for c in "⊕⊗≤≥≠∈∥∧∨→∀√∑⋅×≡≈≅∅∞πσαβγδθλμφψωℝℤ"
       if full.count(c)}
for c, n in sorted(uni.items(), key=lambda x: -x[1]):
    print(f"  {c}  x{n}")

print()
print("=== VI DU CONG THUC DA RENDER (nhung chua bien dich) ===")
shown = 0
for t in texts:
    if len(t) < 20:
        continue
    if any(ch in t for ch in "⊕≤≥≠∈∥∧∨→∀√∑⋅≡≅") and shown < 12:
        print(f"  - {t[:150]}")
        shown += 1

print()
print("=== CONG THUC HIEN THI (render thanh khoi) ===")
import docx
d = docx.Document(str(DOCX))
cnt = 0
for tb in d.tables:
    if len(tb.rows) == 1 and len(tb.columns) == 1:
        txt = tb.cell(0, 0).text.strip()
        if any(ch in txt for ch in "⊕≤≥≠∈∥∧∨→∀√∑⋅≡≅") and len(txt) < 700:
            fonts = {r.font.name for p in tb.cell(0, 0).paragraphs for r in p.runs}
            if "Cambria Math" in fonts:
                print(f"  [{cnt+1}] {txt[:130]}")
                cnt += 1
                if cnt >= 14:
                    break
print(f"\nTong cong thuc hien thi (font Cambria Math): {cnt}")