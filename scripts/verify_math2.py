#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Soi 13 dau $ con sot + dem khoi cong thuc that su (khong phai callout)."""
import re
import sys
import zipfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
DOCX = ROOT / "CloakShare_Bao_Cao_Chuyen_Khao_400_Trang.docx"

with zipfile.ZipFile(DOCX) as z:
    xml = z.read("word/document.xml").decode("utf-8")

texts = [re.sub(r"<[^>]+>", "", t) for t in
         re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, re.DOTALL)]

print("=== DOAN CHUA KY TU '$' ===")
for t in texts:
    if "$" in t:
        i = t.index("$")
        print(f"  ...{t[max(0,i-70):i+70]}...")
        print(f"      full: {t[:180]}")

print()
print("=== KHOI CONG THUC THAT SU (khong phai callout 'GHI CHU:') ===")
import docx
d = docx.Document(str(DOCX))
n = 0
for tb in d.tables:
    if len(tb.rows) == 1 and len(tb.columns) == 1:
        txt = tb.cell(0, 0).text.strip()
        fonts = {r.font.name for p in tb.cell(0, 0).paragraphs for r in p.runs}
        if "Cambria Math" in fonts and not txt.startswith("GHI CHÚ:"):
            n += 1
            if n <= 20:
                print(f"  [{n:>2}] {txt[:120]}")
print(f"\nTong khoi cong thuc hien thi that su: {n}")